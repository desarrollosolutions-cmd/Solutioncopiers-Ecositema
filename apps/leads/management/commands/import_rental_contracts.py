"""
Importa contratos de alquiler desde los Excel históricos (RENTAS ADIS,
RENTAS MARIANA). Normaliza datos sucios/inconsistentes, relaciona contra
clientes (Lead) y equipos (Copier/CopierUnit) ya existentes en vez de
duplicar, y deja todo lo ambiguo anotado en el contrato para revisión
manual -- nunca se descarta información en silencio.

Uso:
    python manage.py import_rental_contracts                  # dry-run, solo reporte
    python manage.py import_rental_contracts --commit          # escribe en la BD
    python manage.py import_rental_contracts --commit --report-file out.txt

Es seguro volver a correrlo: cada contrato importado queda marcado con un
identificador de fila de origen en las notas, y si ya existe uno con ese
identificador no se vuelve a crear.
"""
from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from difflib import SequenceMatcher

from django.core.management.base import BaseCommand
from django.utils import timezone

ADIS_PATH = r"C:\Users\User\Desktop\solution-copiers\RENTAS ADIS (1).xlsx"
MARIANA_PATH = r"C:\Users\User\Desktop\solution-copiers\RENTAS MARIANA 28 AGOSTO (1).xlsx"

MONTHS_ES = {
    "enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6,
    "julio": 7, "agosto": 8, "septiembre": 9, "setiembre": 9, "octubre": 10,
    "noviembre": 11, "diciembre": 12,
}

SITE_QUALIFIER_RE = re.compile(
    r"\b(MAQUINA\s*\d*|PISO\s*\d*|PRIMER\s+PISO|SEGUNDO\s+PISO|TERCER\s+PISO|"
    r"RECEPCION|SERVICIO\s+T[EÉ]CNICO|PUNTO\s+CLAVE.*|SEDE.*|LOCAL\s*\d*)\b",
    re.IGNORECASE,
)
LEGAL_SUFFIX_WORDS = {"SAS", "S", "A", "LTDA", "EU", "SA", "CIA", "Y"}
MODEL_CODE_RE = re.compile(r"[A-Z]{1,4}\s?C?\s?\d{2,5}[A-Z]{0,3}")
PHONE_RE = re.compile(r"(\+?\d[\d\s().+-]{6,17}\d)")


def strip_accents(s: str) -> str:
    return unicodedata.normalize("NFKD", s or "").encode("ascii", "ignore").decode()


def normalize_name(s: str) -> str:
    s = strip_accents(s or "").upper()
    s = SITE_QUALIFIER_RE.sub(" ", s)
    s = re.sub(r"[^A-Z0-9 ]", " ", s)
    words = [w for w in s.split() if w not in LEGAL_SUFFIX_WORDS]
    return " ".join(words)


# Primeras palabras demasiado genéricas como para asumir "mismo cliente" solo
# porque dos nombres las comparten (ej. "CLINICA SANTA MARIA" vs "CLINICA
# SANTA ROSA" NO son el mismo cliente aunque compartan "CLINICA").
GENERIC_LEAD_WORDS = {
    "CLINICA", "CENTRO", "GRUPO", "CORPORACION", "FUNDACION", "COLEGIO",
    "INSTITUTO", "ASOCIACION", "COMPANIA", "EMPRESA", "COOPERATIVA",
    "ORGANIZACION", "COMERCIALIZADORA", "DISTRIBUIDORA", "INVERSIONES",
}


def name_similarity(a_norm: str, b_norm: str) -> float:
    """Ratio de similitud, con un empujón cuando un nombre es claramente el
    "cliente raíz" del otro (ej. "MEDIC COLOMBIA" vs "MEDIC COLOMBIA PISO 6"
    -- mismo cliente con varias sedes, algo muy común en estos datos)."""
    ratio = SequenceMatcher(None, a_norm, b_norm).ratio()
    a_words, b_words = a_norm.split(), b_norm.split()
    if not a_words or not b_words:
        return ratio
    shorter, longer = (a_words, b_words) if len(a_words) <= len(b_words) else (b_words, a_words)
    if len(shorter) >= 2 and shorter[0] not in GENERIC_LEAD_WORDS and longer[: len(shorter)] == shorter:
        ratio = max(ratio, 0.92)
    return ratio


def parse_money(v) -> Decimal | None:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return Decimal(str(v)).quantize(Decimal("0.01"))
    s = str(v).strip()
    if not s:
        return None
    s = re.sub(r"[^\d,.]", "", s)
    if not s:
        return None
    # "1.449.000" o "1,449,000" -> separador de miles; "1449000,50" -> decimal con coma
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif s.count(",") == 1 and len(s.split(",")[1]) == 2:
        s = s.replace(".", "").replace(",", ".")
    else:
        s = s.replace(",", "").replace(".", "")
    try:
        return Decimal(s).quantize(Decimal("0.01"))
    except InvalidOperation:
        return None


def parse_meter(v):
    """Devuelve (valor_int_o_None, sospechoso_bool). Los contadores de página
    son enteros grandes; un valor tipo 140.975 leído como float (=140.975)
    casi seguro perdió los separadores de miles originales -- se marca como
    sospechoso en vez de usarse tal cual."""
    if v is None:
        return None, False
    if isinstance(v, float):
        if not v.is_integer():
            return None, True
        return int(v), False
    if isinstance(v, int):
        return v, False
    s = str(v).strip()
    if re.fullmatch(r"[\d.,]+", s):
        digits = re.sub(r"[.,]", "", s)
        if digits.isdigit():
            return int(digits), False
    return None, True


def parse_copies_plan(v):
    """Extrae (copias_incluidas, precio_copia_extra) de texto libre tipo
    "3500/48", "5.000 / 50", "12000 copias/ copia adici 34 iva incluido".
    Si no se puede interpretar con confianza, devuelve (None, None) y el
    texto original se conserva completo en las notas del contrato."""
    if v is None:
        return None, None
    if isinstance(v, (int, float)):
        return int(v), None
    s = str(v).strip()
    m = re.match(r"^\s*([\d.,]+)\s*/\s*\$?\s*([\d.,]+)", s)
    if m:
        included = re.sub(r"[.,]", "", m.group(1))
        extra = m.group(2).replace(",", ".")
        try:
            return int(included), Decimal(extra).quantize(Decimal("0.0001"))
        except InvalidOperation:
            return None, None
    return None, None


def parse_start_date(v):
    if v is None:
        return None
    if isinstance(v, (date, datetime)):
        return v.date() if isinstance(v, datetime) else v
    s = strip_accents(str(v)).upper()
    m = re.search(r"(\d{1,2})\s*(?:DE)?\s*([A-Z]+)\s*(?:DEL?)?\s*(\d{4})", s)
    if m:
        day = int(m.group(1))
        month = MONTHS_ES.get(m.group(2).lower())
        year = int(m.group(3))
        if month:
            try:
                return date(year, month, day)
            except ValueError:
                return None
    return None


def parse_contact(v):
    """Devuelve (nombre_contacto, telefono) a partir de texto libre tipo
    "LINA GONZALEZ 300 2786849" o "VIVIAN 301 5999148 - Erika 302 4137393".
    Se queda con el primer teléfono encontrado; el texto completo original
    siempre se conserva aparte (notas del lead/contrato)."""
    if not v:
        return "", ""
    s = str(v).strip()
    m = PHONE_RE.search(s)
    phone = re.sub(r"[^\d+]", "", m.group(1)) if m else ""
    name = (s[: m.start()] + s[m.end():]).strip(" /-\n") if m else s
    return name, phone


def extract_current_model(equipo: str) -> str:
    """Si el texto de equipo describe un cambio ("MP301 cambio a IM 350"),
    se queda con el ÚLTIMO código de modelo mencionado (el vigente)."""
    if not equipo:
        return ""
    codes = MODEL_CODE_RE.findall(equipo.upper())
    return codes[-1].strip() if codes else equipo.strip()[:80]


def normalize_model_code(s: str) -> str:
    return re.sub(r"[^A-Z0-9]", "", strip_accents(s or "").upper())


@dataclass
class RawRow:
    source: str
    row_ix: int
    excel_row: int  # fila física en la hoja -- única siempre, a diferencia de row_ix
                     # (la columna de índice del Excel), que puede venir en blanco y
                     # repetirse en varias filas distintas.
    cliente: str
    equipo: str
    copies_plan_raw: str | object
    canon_raw: object
    fecha_corte: str
    meters: list = field(default_factory=list)  # [(label, raw_value)]
    serial: str = ""
    contacto: str = ""
    direccion: str = ""
    fecha_inicio_raw: object = None
    contador_inicial_raw: object = None


class _PendingLead:
    """Marcador liviano para un Lead que SE CREARÍA en modo dry-run (sin
    tocar la BD), para poder simular las coincidencias dentro de la misma
    corrida igual que en modo --commit."""
    __slots__ = ("pk", "full_name")

    def __init__(self, pk, full_name):
        self.pk = pk
        self.full_name = full_name


class _PendingCopier:
    """Igual que _PendingLead pero para fichas de catálogo (Copier) que se
    crearían en dry-run -- así dos filas con el mismo modelo escrito distinto
    (ej. "IM 350" y "IM350") se reconocen como la misma ficha nueva también
    en el reporte de vista previa, no solo al aplicar con --commit."""
    __slots__ = ("pk", "model_number")

    def __init__(self, pk, model_number):
        self.pk = pk
        self.model_number = model_number


class Command(BaseCommand):
    help = "Importa contratos de alquiler desde los Excel de rentas (ADIS y MARIANA)."

    def add_arguments(self, parser):
        parser.add_argument("--commit", action="store_true", help="Escribir en la base de datos (por defecto solo reporta).")
        parser.add_argument("--report-file", default="", help="Ruta donde guardar el reporte detallado (texto).")
        parser.add_argument("--limit", type=int, default=0, help="Procesar solo las primeras N filas (para probar --commit en chico).")
        parser.add_argument("--only-excel-row", type=int, default=0,
                             help="Procesar solo una fila puntual por su número físico en la hoja (ver --only-source).")
        parser.add_argument("--only-source", default="", choices=["", "ADIS", "MARIANA"],
                             help="Combinar con --only-excel-row para indicar de cuál archivo.")
        parser.add_argument("--adis", default=ADIS_PATH)
        parser.add_argument("--mariana", default=MARIANA_PATH)

    # ------------------------------------------------------------------
    def handle(self, *args, **options):
        import openpyxl

        commit = options["commit"]
        rows: list[RawRow] = []

        wb1 = openpyxl.load_workbook(options["adis"], data_only=True)
        ws1 = wb1["Hoja1"]
        for excel_row, r in enumerate(ws1.iter_rows(min_row=3, max_row=ws1.max_row, values_only=True), start=3):
            if not r[1]:
                continue
            rows.append(RawRow(
                source="ADIS", row_ix=r[0] or 0, excel_row=excel_row, cliente=str(r[1]).strip(),
                equipo=str(r[2] or "").strip(), copies_plan_raw=r[3], canon_raw=r[4],
                fecha_corte=str(r[5] or "").strip(),
                meters=[("junio", r[6]), ("julio", r[7]), ("agosto", r[8])],
                serial=str(r[9] or "").strip(), contacto=str(r[10] or "").strip(),
                direccion=str(r[11] or "").strip(), fecha_inicio_raw=r[12],
                contador_inicial_raw=r[13],
            ))

        wb2 = openpyxl.load_workbook(options["mariana"], data_only=True)
        ws2 = wb2["CANON RENTAS"]
        for excel_row, r in enumerate(ws2.iter_rows(min_row=4, max_row=ws2.max_row, values_only=True), start=4):
            if not r[1]:
                continue
            rows.append(RawRow(
                source="MARIANA", row_ix=r[0] or 0, excel_row=excel_row, cliente=str(r[1]).strip(),
                equipo=str(r[2] or "").strip(), copies_plan_raw=r[3], canon_raw=r[4],
                fecha_corte=str(r[8] or "").strip(),
                meters=[("junio", r[5]), ("julio", r[6]), ("agosto", r[7])],
                serial=str(r[11] or "").strip(), contacto=str(r[12] or "").strip(),
                direccion="", fecha_inicio_raw=None, contador_inicial_raw=None,
            ))

        self.stdout.write(f"Filas leídas: ADIS + MARIANA = {len(rows)}")
        if options["limit"]:
            rows = rows[: options["limit"]]
            self.stdout.write(f"--limit activo: procesando solo {len(rows)} filas")
        if options["only_excel_row"]:
            rows = [r for r in rows if r.excel_row == options["only_excel_row"]
                    and (not options["only_source"] or r.source == options["only_source"])]
            self.stdout.write(f"--only-excel-row activo: {len(rows)} fila(s) coinciden")

        from apps.leads.models import Lead, RentalContract
        from apps.catalog.models import Copier, CopierCategory

        all_leads = list(Lead.objects.all().only("id", "full_name", "phone", "address"))
        lead_norm_cache = {lead.pk: normalize_name(lead.full_name) for lead in all_leads}
        self._pending_pk_counter = 0
        all_copiers = list(Copier.objects.all())
        copier_norm_cache = {c.pk: normalize_model_code(c.model_number) for c in all_copiers}

        cat_bn = CopierCategory.objects.filter(slug="fotocopiadoras-blanco-y-negro").first() \
            or CopierCategory.objects.first()
        cat_color = CopierCategory.objects.filter(slug="fotocopiadoras-a-color").first() or cat_bn

        stats = {"leads_matched": 0, "leads_created": 0, "leads_flagged": 0,
                  "copiers_matched": 0, "copiers_created": 0,
                  "units_created": 0, "units_reused": 0,
                  "contracts_created": 0, "contracts_skipped_existing": 0,
                  "rows_flagged": 0}
        report_lines = []
        next_num = self._next_contract_seq(commit)

        for row in rows:
            flags = []
            source_tag = f"[Importado: {row.source} fila-excel {row.excel_row} (índice {row.row_ix})]"

            # Evita duplicar si el comando ya corrió antes con --commit
            if commit and RentalContract.objects.filter(notes__contains=source_tag).exists():
                stats["contracts_skipped_existing"] += 1
                continue

            # ---- Cliente ----
            # Se compara tanto contra los leads ya existentes en la BD como
            # contra los que ya se crearon/simularon en ESTA misma corrida,
            # para que variantes de sede del mismo cliente (ej. "MEDIC
            # COLOMBIA PISO 6", "MEDIC COLOMBIA GUAYABAL"...) queden bajo un
            # solo cliente en vez de crear uno nuevo por cada fila.
            search_name = row.cliente
            norm = normalize_name(search_name)
            best_lead, best_ratio = None, 0.0
            for lead in all_leads:
                ratio = name_similarity(norm, lead_norm_cache[lead.pk])
                if ratio > best_ratio:
                    best_ratio, best_lead = ratio, lead

            contact_name, phone = parse_contact(row.contacto)
            if best_ratio >= 0.90:
                lead = best_lead
                stats["leads_matched"] += 1
                lead_note = f"Coincidencia ({best_ratio:.0%}) con: {lead.full_name}"
            elif best_ratio >= 0.60:
                flags.append(f"Posible cliente duplicado no vinculado automáticamente: '{row.cliente}' se parece "
                             f"({best_ratio:.0%}) a '{best_lead.full_name}' (id {best_lead.pk}) -- revisar y fusionar a mano si aplica.")
                stats["leads_flagged"] += 1
                lead = None
            else:
                lead = None

            if lead is None:
                city = "Medellín"
                addr_up = strip_accents(row.direccion).upper()
                for c in ["CARTAGENA", "BARRANQUILLA", "BOGOTA", "PEREIRA", "CALI", "RIOSUCIO"]:
                    if c in addr_up or c in strip_accents(row.cliente).upper():
                        city = c.title()
                        break
                if commit:
                    lead = Lead.objects.create(
                        full_name=row.cliente[:150],
                        phone=phone[:30],
                        address=row.direccion[:255],
                        city=city,
                        source=Lead.Source.IMPORT,
                        notes_internal=f"Contacto original: {row.contacto}".strip(),
                    )
                    stats["leads_created"] += 1
                else:
                    lead = _PendingLead(pk=self._next_pending_pk(), full_name=row.cliente[:150])
                    stats["leads_created"] += 1
                all_leads.append(lead)
                lead_norm_cache[lead.pk] = normalize_name(lead.full_name)
                lead_note = "Cliente nuevo" + ("" if commit else " [dry-run]")
            elif commit:
                update_fields = []
                if not lead.phone and phone:
                    lead.phone = phone[:30]
                    update_fields.append("phone")
                if not lead.address and row.direccion:
                    lead.address = row.direccion[:255]
                    update_fields.append("address")
                if update_fields:
                    lead.save(update_fields=update_fields)

            # ---- Equipo / catálogo ----
            current_model_text = extract_current_model(row.equipo)
            norm_model = normalize_model_code(current_model_text)
            copier = None
            if norm_model:
                for c in all_copiers:
                    cn = copier_norm_cache[c.pk]
                    if cn and (norm_model == cn or norm_model in cn or cn in norm_model):
                        copier = c
                        break
            copier_is_new = False
            if copier:
                stats["copiers_matched"] += 1
            elif norm_model:
                copier_is_new = True
                model_number = current_model_text[:80] or row.equipo[:80]
                if commit:
                    is_color = "C" in norm_model[:6] and not norm_model.startswith("IM4")
                    category = cat_color if is_color else cat_bn
                    copier = Copier.objects.create(
                        name=current_model_text or row.equipo[:160],
                        brand="Ricoh",
                        model_number=model_number,
                        category=category,
                        short_description="Ficha creada automáticamente por importación de contratos -- completar.",
                        description=f"Equipo importado desde hoja de rentas ({row.source}). "
                                    f"Completar ficha técnica real antes de publicar.",
                        speed_ppm=30,
                        available_for_rental=True,
                    )
                else:
                    copier = _PendingCopier(pk=self._next_pending_pk(), model_number=model_number)
                all_copiers.append(copier)
                copier_norm_cache[copier.pk] = normalize_model_code(model_number)
                stats["copiers_created"] += 1
                flags.append(f"Equipo '{current_model_text}' no existía en el catálogo -- "
                              f"{'se creó' if commit else '[dry-run] se crearía'} una ficha borrador "
                              f"(sin publicar) para poder asociar la unidad. Completar datos técnicos reales.")

            # ---- Unidad física (serial) ----
            unit = None
            serial_clean = re.sub(r"\s+", "", row.serial.split("/")[0].split("CAMBIO")[0].split("cambio")[0]) if row.serial else ""
            if serial_clean and copier and commit:
                from apps.catalog.models import CopierUnit
                unit, was_created = CopierUnit.objects.get_or_create(
                    serial_number=serial_clean,
                    defaults={"copier": copier, "status": CopierUnit.UnitStatus.IN_FIELD},
                )
                if was_created:
                    stats["units_created"] += 1
                else:
                    stats["units_reused"] += 1
                    flags.append(f"Serial '{serial_clean}' ya existía como unidad -- se reutilizó (revisar si es correcto).")

            # ---- Contador ----
            meter_val, meter_flag = None, False
            meter_notes = []
            for label, raw in row.meters:
                v, suspicious = parse_meter(raw)
                meter_notes.append(f"{label}: {raw}")
                if v is not None and not suspicious:
                    meter_val = v
                elif suspicious and raw is not None:
                    meter_flag = True
            if meter_flag:
                flags.append(f"Contador con formato dudoso, revisar valores originales: {'; '.join(meter_notes)}")
            if unit and meter_val and commit:
                unit.current_meter = meter_val
                unit.save(update_fields=["current_meter"])

            # ---- Plan de copias / canon ----
            copies_plan_text = strip_accents(str(row.copies_plan_raw or "")).upper()
            pay_per_page = "LO QUE IMPRIM" in copies_plan_text or "SOLO LO QUE" in copies_plan_text
            if pay_per_page:
                # Sin cuota fija: el valor en CANON es en realidad el precio
                # por copia, no una mensualidad -- guardarlo tal cual como
                # monthly_rate confundiría los reportes financieros.
                copies_included = 0
                overage_rate = parse_money(row.canon_raw)
                monthly_rate = Decimal("0.00")
                flags.append(f"Sin cuota fija -- facturación 100% por página impresa (${overage_rate}/copia). "
                             f"monthly_rate queda en $0 a propósito.")
            else:
                copies_included, overage_rate = parse_copies_plan(row.copies_plan_raw)
                if copies_included is None and row.copies_plan_raw:
                    flags.append(f"Plan de copias en texto libre, no se pudo separar incluidas/excedente: "
                                 f"'{row.copies_plan_raw}' -- queda completo en las notas del contrato.")
                monthly_rate = parse_money(row.canon_raw)
                if monthly_rate is None:
                    flags.append(f"Sin canon/valor mensual legible (valor original: {row.canon_raw!r}) -- se deja en $0, revisar y corregir.")
                    monthly_rate = Decimal("0.00")
                elif monthly_rate < 1000:
                    flags.append(f"Canon inusualmente bajo (${monthly_rate}) para ser una cuota mensual -- revisar si en realidad es "
                                 f"un precio por copia u otro dato mal ubicado en el Excel.")

            # ---- Fecha de inicio ----
            start_date = parse_start_date(row.fecha_inicio_raw)
            if start_date is None:
                start_date = timezone.localdate()
                flags.append(f"Sin fecha de inicio de contrato legible (valor original: {row.fecha_inicio_raw!r}) "
                             f"-- se usó la fecha de hoy como aproximación, corregir manualmente.")

            # ---- Armar notas del contrato (todo lo que no tiene campo propio) ----
            notes_parts = [source_tag]
            if row.equipo and row.equipo != current_model_text:
                notes_parts.append(f"Descripción original de equipo: {row.equipo}")
            if row.fecha_corte:
                notes_parts.append(f"Corte de facturación: {row.fecha_corte}")
            if meter_notes:
                notes_parts.append("Contadores registrados: " + "; ".join(meter_notes))
            if row.contador_inicial_raw:
                notes_parts.append(f"Contador inicial: {row.contador_inicial_raw}")
            if isinstance(row.copies_plan_raw, str) and copies_included is None:
                notes_parts.append(f"Plan de copias (texto original): {row.copies_plan_raw}")
            if row.contacto:
                notes_parts.append(f"Contacto: {row.contacto}")
            if flags:
                notes_parts.append("⚠ REVISAR: " + " | ".join(flags))

            equipment_description = (row.equipo or current_model_text or "Equipo sin descripción")[:300]
            if row.serial:
                equipment_description = f"{equipment_description} — Serial: {row.serial}"[:300]

            report_lines.append(
                f"[{row.source} #{row.row_ix}] {row.cliente!r} -> lead={'match' if best_ratio>=0.90 else ('flag' if best_ratio>=0.60 else 'nuevo')} "
                f"({best_ratio:.0%}) | equipo={current_model_text!r} catálogo={'nuevo' if copier_is_new else ('sí' if copier else 'no')} | "
                f"canon={monthly_rate} | inicio={start_date} | flags={len(flags)}"
            )
            if flags:
                stats["rows_flagged"] += 1

            if commit:
                contract = RentalContract.objects.create(
                    lead=lead,
                    copier=copier,
                    unit=unit,
                    equipment_description=equipment_description,
                    contract_number=f"IMP-{next_num:04d}",
                    status=RentalContract.Status.ACTIVE,
                    start_date=start_date,
                    monthly_rate=monthly_rate,
                    copies_included=copies_included or 0,
                    copy_overage_rate=overage_rate,
                    notes="\n".join(notes_parts),
                )
                next_num += 1
                stats["contracts_created"] += 1

        # ---- Reporte ----
        self.stdout.write("\n".join(report_lines))
        self.stdout.write("\n" + "=" * 60)
        for k, v in stats.items():
            self.stdout.write(f"{k}: {v}")
        self.stdout.write("=" * 60)
        if not commit:
            self.stdout.write(self.style.WARNING("DRY-RUN: no se escribió nada en la base de datos. Usa --commit para aplicar."))
        else:
            self.stdout.write(self.style.SUCCESS("Importación aplicada."))

        report_file = options["report_file"]
        if report_file:
            with open(report_file, "w", encoding="utf-8") as f:
                f.write("\n".join(report_lines))
                f.write("\n\n" + "\n".join(f"{k}: {v}" for k, v in stats.items()))

    def _next_contract_seq(self, commit: bool) -> int:
        from apps.leads.models import RentalContract
        existing = RentalContract.objects.filter(contract_number__startswith="IMP-").count()
        return existing + 1

    def _next_pending_pk(self) -> int:
        """PKs sintéticos negativos para leads simulados en dry-run, para que
        nunca choquen con un pk real de la base de datos."""
        self._pending_pk_counter -= 1
        return self._pending_pk_counter
